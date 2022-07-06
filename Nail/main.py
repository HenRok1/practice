import openpyxl
from docx import Document


def replace_word(sheet):
    for type_name in range(2, sheet.max_row + 1):
        table_key = sheet[type_name][2].value
        if table_key in cell.text:
            table_name = sheet[type_name][1].value
            quest = input("Введите " + table_name + ": ")
            cell.text = cell.text.replace(table_key, quest)


path_1 = input("Введите путь к файлу .docx: ")
path_2 = input("Введите путь к файлу .xlsx: ")

wordDoc = Document(path_1)
book = openpyxl.open(path_2, read_only=True)
sheet = book.worksheets[0]


tables = wordDoc.tables
table = tables[0]

for i, row in enumerate(table.rows):
    for j, cell in enumerate(row.cells):
        if cell.text:
            if '${' in cell.text:
                replace_word(sheet)

wordDoc.save("example.docx")